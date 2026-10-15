import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


class Note:
    def __init__(self, matiere, valeur):
        self.matiere = matiere
        self.valeur = valeur

    def afficher(self):
        print(f"{self.matiere} : {self.valeur}/20")

import statistics

class Etudiant:
    def __init__(self, nom, prenom, age, filiere, niveau=""):
        self.nom = nom
        self.prenom = prenom
        self.age = age
        self.filiere = filiere
        self.niveau = niveau
        self.notes = []

    def ajouter_note(self, note):
        self.notes.append(note)

    def supprimer_note(self, matiere):
        self.notes = [n for n in self.notes if n.matiere != matiere]

    def calculer_moyenne(self):
        if len(self.notes) == 0:
            return 0
        return statistics.mean([n.valeur for n in self.notes])

    def afficher(self):
        print(f"👤 {self.nom} {self.prenom} | Age: {self.age} | Filière: {self.filiere} | Niveau: {self.niveau}")
        print(f"   Moyenne : {self.calculer_moyenne():.2f}/20")
        for n in self.notes:
            n.afficher()


class Promotion:
    def __init__(self, filiere, niveau):
        self.filiere = filiere
        self.niveau = niveau
        self.etudiants = []

    @property
    def nom_promotion(self):
        return f"{self.filiere} {self.niveau}".strip()

    def ajouter_etudiant(self, etudiant):
        self.etudiants.append(etudiant)
        print(f"✅ {etudiant.nom} {etudiant.prenom} ajouté avec succès !")

    def supprimer_etudiant(self, nom, prenom):
        self.etudiants = [e for e in self.etudiants if not (e.nom == nom and e.prenom == prenom)]
        print(f"🗑️ Etudiant {nom} {prenom} supprimé !")

    def rechercher_etudiant(self, nom, prenom):
        for e in self.etudiants:
            if e.nom == nom and e.prenom == prenom:
                return e
        return None

    def meilleur_etudiant(self):
        if not self.etudiants:
            return None
        return max(self.etudiants, key=lambda e: e.calculer_moyenne())

    def moyenne_generale(self):
        if not self.etudiants:
            return 0
        return statistics.mean([e.calculer_moyenne() for e in self.etudiants])

    def afficher_tous(self):
        print(f"\n📋 Promotion : {self.nom_promotion}")
        print(f"   Nombre d'étudiants : {len(self.etudiants)}")
        print(f"   Moyenne générale   : {self.moyenne_generale():.2f}/20")
        print("-" * 40)
        for e in self.etudiants:
            e.afficher()
            print("-" * 40)


class Etablissement:
    def __init__(self, nom="Établissement"):
        self.nom = nom
        self.promotions = {}  # (filiere, niveau) -> Promotion

    def get_or_create(self, filiere, niveau):
        key = (filiere, niveau)
        if key not in self.promotions:
            self.promotions[key] = Promotion(filiere, niveau)
        return self.promotions[key]

    def get(self, filiere, niveau):
        return self.promotions.get((filiere, niveau))

    def supprimer(self, filiere, niveau):
        self.promotions.pop((filiere, niveau), None)

    def liste_promotions(self):
        return list(self.promotions.values())


import csv

class GestionFichier:

    # ── CSV ──────────────────────────────────────
    def exporter_csv(self, etudiants, nom_fichier):
        with open(nom_fichier, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(["Nom", "Prénom", "Age", "Filière", "Niveau", "Matière", "Note"])
            for e in etudiants:
                if e.notes:
                    for n in e.notes:
                        writer.writerow([e.nom, e.prenom, e.age, e.filiere, e.niveau, n.matiere, n.valeur])
                else:
                    writer.writerow([e.nom, e.prenom, e.age, e.filiere, e.niveau, "", ""])
        print(f"✅ CSV exporté : {nom_fichier}")

    def importer_csv(self, nom_fichier):
        etudiants = {}
        with open(nom_fichier, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                cle = f"{row['Nom']}_{row['Prénom']}"
                if cle not in etudiants:
                    etudiants[cle] = Etudiant(
                        row["Nom"], row["Prénom"], int(row["Age"]),
                        row["Filière"], row.get("Niveau", ""),
                    )
                if row.get("Matière"):
                    etudiants[cle].ajouter_note(Note(row["Matière"], float(row["Note"])))
        print(f"✅ CSV importé : {nom_fichier}")
        return list(etudiants.values())

    # ── EXCEL ────────────────────────────────────
    def exporter_excel(self, etudiants, nom_fichier):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = Workbook()
        wb.remove(wb.active)

        title_font = Font(size=14, bold=True, color="FFFFFF")
        header_font = Font(bold=True, color="FFFFFF")
        title_fill = PatternFill("solid", fgColor="2C3E50")
        header_fill = PatternFill("solid", fgColor="3498DB")
        moyenne_fill = PatternFill("solid", fgColor="2ECC71")
        center = Alignment(horizontal="center", vertical="center")

        for e in etudiants:
            base = f"{e.nom}_{e.prenom}"[:31]
            sheet_name, n = base, 1
            while sheet_name in wb.sheetnames:
                n += 1
                sheet_name = f"{base[:28]}_{n}"
            ws = wb.create_sheet(sheet_name)

            # Bandeau titre
            ws.merge_cells("A1:B1")
            ws["A1"] = f"{e.nom} {e.prenom}"
            ws["A1"].font = title_font
            ws["A1"].fill = title_fill
            ws["A1"].alignment = center
            ws.row_dimensions[1].height = 28

            # Infos étudiant
            for i, (label, value) in enumerate(
                [("Filière", e.filiere), ("Niveau", e.niveau), ("Age", e.age)],
                start=2,
            ):
                ws[f"A{i}"] = label
                ws[f"A{i}"].font = Font(bold=True)
                ws[f"B{i}"] = value

            # En-tête tableau notes
            HEADER_ROW = 6
            ws[f"A{HEADER_ROW}"] = "Matière"
            ws[f"B{HEADER_ROW}"] = "Note /20"
            for cell in (ws[f"A{HEADER_ROW}"], ws[f"B{HEADER_ROW}"]):
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center

            # Notes
            row = HEADER_ROW + 1
            for note in e.notes:
                ws[f"A{row}"] = note.matiere
                ws[f"B{row}"] = note.valeur
                row += 1

            # Ligne Moyenne
            ws[f"A{row}"] = "Moyenne"
            ws[f"B{row}"] = round(e.calculer_moyenne(), 2)
            for cell in (ws[f"A{row}"], ws[f"B{row}"]):
                cell.font = Font(bold=True)
                cell.fill = moyenne_fill

            ws.column_dimensions["A"].width = 28
            ws.column_dimensions["B"].width = 16

        wb.save(nom_fichier)
        print(f"✅ Excel exporté : {nom_fichier}")

    def importer_excel(self, nom_fichier):
        from openpyxl import load_workbook

        wb = load_workbook(nom_fichier, data_only=True)
        etudiants = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # Récupère le nom depuis le bandeau (A1) ou retombe sur le nom d'onglet
            titre = ws["A1"].value or sheet_name
            parts = str(titre).strip().split(" ", 1)
            nom = parts[0]
            prenom = parts[1] if len(parts) > 1 else ""

            filiere = ws["B2"].value or ""
            niveau = ws["B3"].value or ""
            try:
                age = int(ws["B4"].value or 0)
            except (TypeError, ValueError):
                age = 0

            e = Etudiant(str(nom), str(prenom), age, str(filiere), str(niveau))

            # Parcourt les lignes en cherchant l'en-tête "Matière"
            start = None
            for r in range(1, ws.max_row + 1):
                if str(ws[f"A{r}"].value).strip().lower() == "matière":
                    start = r + 1
                    break
            if start is None:
                continue

            for r in range(start, ws.max_row + 1):
                m = ws[f"A{r}"].value
                v = ws[f"B{r}"].value
                if m is None or str(m).strip().lower() == "moyenne":
                    continue
                try:
                    e.ajouter_note(Note(str(m), float(v)))
                except (TypeError, ValueError):
                    continue
            etudiants.append(e)

        print(f"✅ Excel importé : {nom_fichier}")
        return etudiants

class Admin:
    LOGIN = "admin"
    MOT_DE_PASSE = "1234"

    def authentifier(self, login, mot_de_passe):
        if login == self.LOGIN and mot_de_passe == self.MOT_DE_PASSE:
            print("✅ Authentification réussie !")
            return True
        print("❌ Login ou mot de passe incorrect !")
        return False


def saisir_etudiant(filiere_defaut="", niveau_defaut=""):
    nom = input("Nom : ").strip()
    prenom = input("Prénom : ").strip()
    while True:
        try:
            age = int(input("Age : "))
            break
        except ValueError:
            print("⚠️  Age invalide, entre un nombre entier.")
    return Etudiant(nom, prenom, age, filiere_defaut, niveau_defaut)


def saisir_note():
    matiere = input("Matière : ").strip()
    while True:
        try:
            valeur = float(input("Note /20 : "))
            if 0 <= valeur <= 20:
                return Note(matiere, valeur)
            print("⚠️  La note doit être entre 0 et 20.")
        except ValueError:
            print("⚠️  Valeur invalide.")


def menu_promotion(promo, fichiers):
    options = {
        "1": "Ajouter un étudiant",
        "2": "Supprimer un étudiant",
        "3": "Ajouter une note à un étudiant",
        "4": "Supprimer une note d'un étudiant",
        "5": "Rechercher un étudiant",
        "6": "Afficher tous les étudiants",
        "7": "Meilleur étudiant",
        "8": "Exporter en CSV",
        "9": "Importer depuis CSV",
        "10": "Exporter en Excel",
        "11": "Importer depuis Excel",
        "0": "← Retour à la liste des promotions",
    }
    while True:
        print(f"\n===== {promo.nom_promotion} =====")
        for k, v in options.items():
            print(f" {k}. {v}")
        choix = input("Choix : ").strip()

        if choix == "1":
            promo.ajouter_etudiant(saisir_etudiant(promo.filiere, promo.niveau))
        elif choix == "2":
            nom = input("Nom : ").strip(); prenom = input("Prénom : ").strip()
            promo.supprimer_etudiant(nom, prenom)
        elif choix in ("3", "4", "5"):
            nom = input("Nom : ").strip(); prenom = input("Prénom : ").strip()
            etu = promo.rechercher_etudiant(nom, prenom)
            if not etu:
                print("❌ Étudiant introuvable."); continue
            if choix == "3":
                etu.ajouter_note(saisir_note()); print("✅ Note ajoutée.")
            elif choix == "4":
                m = input("Matière à supprimer : ").strip()
                etu.supprimer_note(m); print("🗑️  Note supprimée.")
            else:
                etu.afficher()
        elif choix == "6":
            promo.afficher_tous()
        elif choix == "7":
            top = promo.meilleur_etudiant()
            if top: print("🏆 Meilleur étudiant :"); top.afficher()
            else: print("Aucun étudiant.")
        elif choix == "8":
            nf = input("Nom du fichier CSV : ").strip() or "etudiants.csv"
            fichiers.exporter_csv(promo.etudiants, nf)
        elif choix == "9":
            nf = input("Nom du fichier CSV : ").strip()
            for e in fichiers.importer_csv(nf):
                e.filiere = promo.filiere; e.niveau = promo.niveau
                promo.ajouter_etudiant(e)
        elif choix == "10":
            nf = input("Nom du fichier Excel : ").strip() or "etudiants.xlsx"
            fichiers.exporter_excel(promo.etudiants, nf)
        elif choix == "11":
            nf = input("Nom du fichier Excel : ").strip()
            for e in fichiers.importer_excel(nf):
                e.filiere = promo.filiere; e.niveau = promo.niveau
                promo.ajouter_etudiant(e)
        elif choix == "0":
            return
        else:
            print("⚠️  Choix invalide.")


def menu():
    admin = Admin()
    print("=== Connexion administrateur ===")
    for _ in range(3):
        login = input("Login : ")
        mdp = input("Mot de passe : ")
        if admin.authentifier(login, mdp):
            break
    else:
        print("🚫 Trop de tentatives. Fin du programme.")
        return

    etab = Etablissement()
    fichiers = GestionFichier()

    while True:
        print("\n========= ÉTABLISSEMENT =========")
        promos = etab.liste_promotions()
        if promos:
            print("Promotions existantes :")
            for i, p in enumerate(promos, 1):
                print(f"  {i}. {p.nom_promotion}  ({len(p.etudiants)} étudiant(s))")
        else:
            print("(aucune promotion)")
        print("\n  N. Nouvelle promotion (filière + niveau)")
        print("  0. Quitter")
        choix = input("Choix : ").strip().lower()

        if choix == "0":
            print("👋 Au revoir !"); break
        elif choix == "n":
            filiere = input("Filière : ").strip()
            niveau = input("Niveau (ex: L1, L2, M1) : ").strip()
            if not filiere or not niveau:
                print("⚠️  Filière et niveau obligatoires."); continue
            promo = etab.get_or_create(filiere, niveau)
            menu_promotion(promo, fichiers)
        else:
            try:
                idx = int(choix) - 1
                if 0 <= idx < len(promos):
                    menu_promotion(promos[idx], fichiers)
                else:
                    print("⚠️  Choix invalide.")
            except ValueError:
                print("⚠️  Choix invalide.")


if __name__ == "__main__":
    menu()
